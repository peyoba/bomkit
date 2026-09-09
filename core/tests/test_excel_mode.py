"""Excel默认模式：旧工具能力不得被网页确认或新模板替换。"""

import io
import json
from pathlib import Path

import pytest
from bomkit_cli.main import _save_with_retry, main
from openpyxl import Workbook, load_workbook

from bomcore.excel_api import analyze_excel, convert_excel
from bomcore.review_api import ReviewSession
from bomcore.schema import ProfileError

from .excel_golden import sheet_signature

CASES = json.loads((Path(__file__).parent / "fixtures/excel_legacy_golden.json").read_text())["cases"]
HEADER = ["Name", "Designator", "Quantity", "Footprint", "Device"]
MATERIAL = [["编码", "名称", "规格型号"], ["001", "电阻", "10kΩ/0603"]]


@pytest.mark.parametrize("case", CASES, ids=lambda c: c["name"])
@pytest.mark.parametrize("trace", [False, True])
def test_old_company_main_is_unchanged(case, trace):
    result = convert_excel({k: case[k] for k in ("bom_rows", "material_rows", "meta")} | {"include_trace": trace})
    wb = load_workbook(io.BytesIO(result["data"]))
    assert sheet_signature(wb.active) == case["expected"]
    assert len(wb.sheetnames) == (3 if trace else 1)
    if trace:
        assert [list(row) for row in wb["原始输入"].iter_rows(values_only=True)] == [
            [v or None for v in row] for row in case["bom_rows"]
        ]
    assert not any(c.data_type == "f" for ws in wb for row in ws for c in row)
    wb.close()


def test_multi_candidates_download_without_web_confirmation():
    case = next(c for c in CASES if c["name"] == "multiple_candidates")
    result = convert_excel({"bom_rows": case["bom_rows"], "material_rows": case["material_rows"]})
    assert result["stats"]["multi"] == 1
    assert result["stats"]["quantity"] == 2
    assert result["stats"]["display_quantity"] == 4
    assert result["stats"]["output_rows"] == 2


def test_grouping_keeps_old_first_nonempty_values_and_trace_preserves_all():
    case = next(c for c in CASES if c["name"] == "same_value_distinct_mpn")
    r = convert_excel({"bom_rows": case["bom_rows"], "include_trace": True})
    w = load_workbook(io.BytesIO(r["data"]))
    assert r["stats"]["groups"] == 1
    assert w.active["H5"].value == "品牌A"
    assert w.active["J5"].value == "PART-A"
    assert "组内型号不同" in w["Excel校对提示"]["C2"].value
    assert w["原始输入"]["E3"].value == "PART-B"


@pytest.mark.parametrize("value,model", [("10kΩ", "10kΩ/0805"), ("10kΩ/0805", "10kΩ/0603")])
def test_package_contradiction_is_reported_in_excel_not_download_blocked(value, model):
    rows = [HEADER, [value, "R1", "1", "0603", model]]
    result = convert_excel({"bom_rows": rows, "material_rows": MATERIAL, "include_trace": True})
    w = load_workbook(io.BytesIO(result["data"]))
    assert "封装互相矛盾" in w["Excel校对提示"]["C2"].value
    assert "0805" in w["Excel校对提示"]["C2"].value
    # 可选v2不能继续错误自动放行；默认Excel导出则继续允许在表格中处理。
    s = ReviewSession(rows, MATERIAL)
    assert s.items[0]["requires_review"]


def test_chinese_ohm_is_known_equivalence_in_optional_review():
    s = ReviewSession([HEADER, ["10k欧", "R1", "1", "0603", "10k欧"]], MATERIAL)
    assert s.items[0]["review_status"] == "auto_passed"


def test_old_non_component_does_not_create_missing_material_task():
    c = next(c for c in CASES if c["name"] == "known_non_material")
    r = convert_excel({"bom_rows": c["bom_rows"], "material_rows": c["material_rows"], "include_trace": True})
    w = load_workbook(io.BytesIO(r["data"]))
    assert r["stats"]["non_component"] == 1
    assert r["stats"]["unmatched"] == 0
    assert w.active["L5"].value == "非物料"
    assert "未匹配" not in str(w["Excel校对提示"]["C2"].value)


def test_5000_grouped_designators_not_silently_truncated():
    rows = [HEADER] + [["10kΩ", "R" + str(n), "1", "0603", "10kΩ/0603"] for n in range(1, 5001)]
    r = convert_excel({"bom_rows": rows, "material_rows": MATERIAL})
    w = load_workbook(io.BytesIO(r["data"]))
    assert r["stats"]["source_rows"] == 5000
    assert r["stats"]["groups"] == r["stats"]["output_rows"] == 1
    assert w.active["F5"].value == 5000
    assert len(w.active["E5"].value.split(",")) == 5000
    assert w.active["E5"].value.endswith("R5000")


@pytest.mark.parametrize("bad", ["", "-1", "NaN", "Infinity", "bad"])
def test_bad_quantity_fails_before_download(bad):
    with pytest.raises(ProfileError):
        convert_excel({"bom_rows": [HEADER, ["10kΩ", "R1", bad, "0603", "X"]]})


def test_material_header_scan_and_legacy_last_duplicate_lookup():
    rows = [["说明"], ["规格", "物料编码", "金蝶系统型号"], ["A", "001", "OLD"], ["A", "002", "NEW"]]
    a = analyze_excel([HEADER, ["A", "U1", "1", "SOIC8", "A"]], rows)
    assert a["items"][0]["match"]["candidates"][0]["code"] == "002"
    assert a["items"][0]["match"]["candidates"][0]["spec"] == "NEW"


def _xlsx(path, rows):
    w = Workbook()
    for row in rows:
        w.active.append(row)
    w.save(path)


def test_cli_keeps_original_flags_and_old_material_format(tmp_path):
    c = next(c for c in CASES if c["name"] == "legacy_simple_material")
    b, m, out = [tmp_path / n for n in ("bom.xlsx", "material.xlsx", "chosen.xlsx")]
    _xlsx(b, c["bom_rows"])
    _xlsx(m, c["material_rows"])
    out.write_bytes(b"old explicit output")
    assert main([str(b), "-m", str(m), "-o", str(out), "--pcba-name", "COMPANY", "--pcb-model", "PCB1"]) == 0
    w = load_workbook(out)
    assert w.active["B5"].value == "001"
    assert w.active["D2"].value == "COMPANY"
    assert w.active["D4"].value == "PCB1"
    assert len(w.sheetnames) == 1


def test_cli_auto_filename_no_clobber_and_optional_trace(tmp_path):
    path = tmp_path / "bom.xlsx"
    _xlsx(path, [HEADER, ["A", "U1", "1", "SOIC8", "A"]])
    sentinel = tmp_path / "bom_converted.xlsx"
    sentinel.write_bytes(b"keep")
    assert main([str(path), "--with-trace"]) == 0
    assert sentinel.read_bytes() == b"keep"
    assert len(load_workbook(tmp_path / "bom_converted_1.xlsx").sheetnames) == 3


def test_cli_permission_error_retries_three_times(monkeypatch, tmp_path):
    import bomkit_cli.main as cli_module

    attempts = []

    def locked(self, data):
        attempts.append(1)
        raise PermissionError("locked")

    monkeypatch.setattr(Path, "write_bytes", locked)
    monkeypatch.setattr(cli_module.time, "sleep", lambda _: None)
    with pytest.raises(PermissionError):
        _save_with_retry(str(tmp_path / "out.xlsx"), b"data")
    assert len(attempts) == 3


def test_meta_and_source_formula_like_values_remain_text():
    r = convert_excel(
        {
            "bom_rows": [HEADER, ["=1+1", "R1", "1", "0603", "=1+1"]],
            "meta": {"pcba_name": "=1+1", "pcb_model": "@danger"},
            "include_trace": True,
        }
    )
    w = load_workbook(io.BytesIO(r["data"]))
    assert w.active["D2"].value == "=1+1"
    assert not any(c.data_type == "f" for ws in w for row in ws for c in row)
