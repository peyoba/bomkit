"""问题驱动校对：确定项零操作，有疑点必须给出可检查的字段证据。"""

import copy
import io

import pytest
from openpyxl import load_workbook

from bomcore.review_api import ReviewSession
from bomcore.review_export import AMBER, export_review
from bomcore.review_rules import rc_value
from bomcore.schema import ProfileError

HEADER = ["Designator", "Quantity", "Name", "Device", "Footprint", "Tolerance", "Manufacturer"]
MAT_HEADER = ["编码", "名称", "规格型号", "禁用状态", "封装", "精度", "厂商"]


def make(
    model="IC-123-A", spec=None, package="SOIC8", library_package="", value=None, tolerance="", brand="", extra=()
):
    bom = [HEADER, ["U1", "1", value if value is not None else model, model, package, tolerance, brand]]
    rc = rc_value(model)
    name = {"R": "合成电阻", "C": "合成电容"}.get(rc[0] if rc else "", "合成测试物料")
    materials = [MAT_HEADER, ["0001", name, spec if spec is not None else model, "否", library_package, "", ""], *extra]
    return ReviewSession(bom, materials)


def codes(item):
    return {p["code"] for p in item["review_findings"]}


@pytest.mark.parametrize(
    "model,spec,package",
    [
        (" IC-123-A ", "IC-123-A", "SOIC8"),
        ("ic-123-a", "IC-123-A/SOIC-8", "SOIC8"),
        ("CAP-123-A", "100nF±10%/50V/C0603(CAP-123-A)", "C0603"),
        ("IC-123-A_C123456", "IC-123-A/SOIC8", "SOIC8"),
        ("IC-123-A", "SOIC8 (IC-123-A)", "SOIC8"),
    ],
)
def test_unique_complete_model_has_no_review_or_manual_audit(model, spec, package):
    s = make(model=model, spec=spec, package=package)
    item = s.items[0]
    assert item["review_findings"] == []
    assert item["review_status"] == "auto_passed"
    assert item["export_ready"] and not item["confirmed"] and not item["confirmation"]
    assert item["original_model"] == model
    assert s.snapshot()["stats"]["pending"] == 0
    s.assert_confirmed()


@pytest.mark.parametrize(
    "model,spec,package",
    [
        ("0.1μF", "100nF±10%/50V/C0603", "C0603"),
        ("0.047uF/50V", "47nF±10%/50V/C0603", "0603C"),
        ("10kΩ", "10000R/1%/R0603", "R1608"),
        ("4R7", "4.7Ω/R0603", "0603"),
        ("0R", "0Ω/R0603", "0603"),
    ],
)
def test_exact_electrical_equivalence_is_not_a_text_difference(model, spec, package):
    s = make(model=model, spec=spec, package=package)
    assert s.items[0]["review_findings"] == []
    assert s.items[0]["export_ready"]


def test_known_parameters_disambiguate_candidates_instead_of_asking_about_all():
    s = make(
        model="100nF",
        spec="100nF/0805",
        package="0603",
        extra=[
            ["0002", "电容", "0.1uF/0603", "否", "", "", ""],
        ],
    )
    item = s.items[0]
    assert item["candidate_count"] == 2 and item["qualified_count"] == 1
    assert item["selected_material"]["code"] == "0002"
    assert item["export_ready"]


def test_two_equally_valid_codes_are_not_auto_selected_by_priority():
    s = make(extra=[["01.9999", "另一个编码", "IC-123-A", "否", "SOIC8", "", ""]])
    assert "ambiguous" in codes(s.items[0])
    assert s.snapshot()["stats"]["pending"] == 1
    with pytest.raises(ProfileError):
        export_review(s, mode="final")


def test_identical_duplicate_library_rows_do_not_create_ambiguity():
    s = make(extra=[["0001", "合成测试物料", "IC-123-A", "否", "", "", ""]])
    assert s.items[0]["qualified_count"] == 1 and s.items[0]["export_ready"]


def test_incomplete_competing_candidate_cannot_be_silently_excluded():
    s = make(
        model="100nF",
        spec="100nF/0603",
        package="0603",
        extra=[
            ["0002", "电容", "100nF", "否", "", "", ""],
        ],
    )
    assert s.items[0]["qualified_count"] == 1
    assert "ambiguous" in codes(s.items[0]) and not s.items[0]["export_ready"]


@pytest.mark.parametrize(
    "model,spec,package",
    [
        ("C0603", "100nF/50V/C0603", "C0603"),
        ("R0603", "10kΩ/R0603", "R0603"),
        ("SOIC8", "IC-123/SOIC8", "SOIC8"),
        ("X7R", "100nF/X7R/0603", "0603"),
        ("50V", "100nF/50V/0603", "0603"),
        ("C0603", "C0603", "C0603"),
        ("C0603_C123456", "100nF/C0603", "C0603"),
    ],
)
def test_package_or_attribute_token_is_not_a_complete_part_identity(model, spec, package):
    s = make(model=model, spec=spec, package=package)
    assert "incomplete_model" in codes(s.items[0])
    assert not s.items[0]["export_ready"]


def test_duplicate_code_with_different_spec_is_a_real_problem():
    s = make(extra=[["0001", "其他料", "IC-123-B", "否", "", "", ""]])
    assert "duplicate_code" in codes(s.items[0])
    assert not s.items[0]["export_ready"]


@pytest.mark.parametrize(
    "kwargs,expected",
    [
        ({"model": "10mΩ", "spec": "10MΩ/0603", "package": "0603"}, "value_mismatch"),
        ({"model": "10mΩ/0603", "spec": "10MΩ/0603", "package": "0603"}, "value_mismatch"),
        ({"model": "10kΩ", "spec": "10.05kΩ/0603", "package": "0603"}, "value_mismatch"),
        ({"model": "100nF", "spec": "100nF/0805", "package": "0603"}, "package_mismatch"),
        ({"spec": "IC-123-A/SOIC16", "package": "SOIC8"}, "package_mismatch"),
        ({"model": "100nF", "spec": "100nF", "package": "0603"}, "missing_package_evidence"),
        ({"model": "100nF/50V", "spec": "100nF/25V/0603", "package": "0603"}, "property_mismatch"),
        ({"model": "100nF/50V/C0G", "spec": "100nF/50V/X7R/0603", "package": "0603"}, "property_mismatch"),
        ({"model": "IC-123-A", "spec": "IC-123-B", "package": "SOIC8"}, "unmatched"),
        ({"model": "IC-123", "spec": "IC-123/NOPB/SOIC8", "package": "SOIC8"}, "model_mismatch"),
        ({"model": "100nF/UNKNOWN", "spec": "100nF/0603", "package": "0603"}, "unverified_parameters"),
        ({"model": "CAP-123", "spec": "100nF/C0603(CAP-123)", "value": "10nF", "package": "0603"}, "value_mismatch"),
    ],
)
def test_uncertain_or_conflicting_items_keep_actionable_evidence(kwargs, expected):
    s = make(**kwargs)
    item = s.items[0]
    assert expected in codes(item), item["review_findings"]
    assert item["requires_review"] and not item["export_ready"]
    for issue in item["review_findings"]:
        assert issue["label"] and issue["reason"]
        assert {"original", "library", "final"}.issubset(issue)


def test_tolerance_and_brand_are_checked_inside_the_library_spec():
    s = make(
        model="CAP-123",
        spec="100nF±10%/50V/C0603(CAP-123)/AA(合成甲厂)",
        package="C0603",
        value="100nF",
        tolerance="±5%",
        brand="BB(合成乙厂)",
    )
    assert {"property_mismatch", "manufacturer_mismatch"}.issubset(codes(s.items[0]))


def test_conflicting_bom_value_columns_are_not_ignored_by_exact_match():
    s = make(model="10kΩ", spec="10kΩ/0603", value="20kΩ", package="0603")
    assert "source_value_conflict" in codes(s.items[0])
    assert not s.items[0]["export_ready"]


def test_property_difference_quotes_the_field_that_actually_contains_it():
    s = make(model="100nF/50V", spec="100nF/25V/0603", value="100nF", package="0603")
    problem = next(p for p in s.items[0]["review_findings"] if p["field"] == "voltage")
    assert "50V" in problem["original"]
    assert "25V" in problem["library"]


def test_source_anomalies_mark_all_involved_rows():
    s = make()
    rows = copy.deepcopy(s.raw_rows)
    rows.append(["u1", "1", "IC-123-A", "IC-123-A", "SOIC8", "", ""])
    duplicate = ReviewSession(rows, [MAT_HEADER, ["0001", "IC", "IC-123-A", "否", "", "", ""]])
    assert all("source_issue" in codes(i) for i in duplicate.items)
    assert duplicate.snapshot()["stats"]["pending"] == 2


def test_only_uncertain_item_blocks_and_confirmation_does_not_touch_auto_item():
    rows = [
        HEADER,
        ["U1", "1", "IC-123-A", "IC-123-A", "SOIC8", "", ""],
        ["U2", "1", "IC-123-A", "IC-123-A", "SOIC16", "", ""],
    ]
    s = ReviewSession(rows, [MAT_HEADER, ["0001", "IC", "IC-123-A/SOIC8", "否", "", "", ""]])
    assert s.snapshot()["stats"] == {"rows": 2, "quantity": 2, "confirmed": 0, "auto_passed": 1, "pending": 1}
    s.confirm(1, "合成测试员")
    s.assert_confirmed()
    assert s.snapshot()["stats"]["confirmed"] == 1
    assert s.items[0]["confirmation"] is None and s.items[0]["history"] == []
    s.update(1, {"final": dict(s.items[1]["final"])})
    assert s.items[1]["confirmed"], "无实际变化的保存不应撤销确认"
    s.update(0, {"final": {"code": "DIFFERENT"}})
    assert not s.items[0]["export_ready"]
    with pytest.raises(ProfileError):
        s.assert_confirmed()
    s.update(0, {"final": {"code": "0001"}})
    s.assert_confirmed()


def test_exports_do_not_mark_clear_rows_or_claim_human_confirmation():
    s = make()
    for mode in ("draft", "final"):
        wb = load_workbook(io.BytesIO(export_review(s, mode=mode)))
        ws = wb.active
        assert ws.cell(3, 8).value is None
        assert ws.cell(3, 5).fill != AMBER
        audit = wb["校对记录"]
        headers = {c.value: c.column for c in audit[1]}
        assert audit.cell(2, headers["状态"]).value == "自动通过（无需人工确认）"
        assert audit.cell(2, headers["校对人"]).value is None
        assert audit.cell(2, headers["确认时间"]).value is None
        assert audit.cell(2, headers["自动判定依据"]).value
        assert audit.cell(2, headers["问题明细"]).value == "[]"
        if mode == "final":
            assert ws.title == "BOM" and "已人工确认" not in ws["A1"].value
