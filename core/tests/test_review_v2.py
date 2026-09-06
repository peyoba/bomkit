"""v2闭环回归仅使用合成数据；真实样例走显式 opt-in 私有回归脚本。"""

import base64
import copy
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from bomcore.review_api import ReviewSession, dispatch, import_materials
from bomcore.review_export import export_review
from bomcore.review_import import detect_input, import_bom, read_file
from bomcore.schema import ProfileError

HEADERS = ["Designator", "Quantity", "Name", "Device", "Footprint", "Manufacturer", "Comment", "Supplier Part"]
BOM = [HEADERS, ["R1", "1", "10kΩ", "R-10K-A", "0603", "", "", "C123456"]]
MATERIAL = [["编码", "名称", "规格型号", "禁用状态", "封装"], ["01.0001", "合成电阻", "R-10K-A", "否", "0603"]]


def session(rows=None, materials=None):
    return ReviewSession(copy.deepcopy(rows or BOM), copy.deepcopy(materials or MATERIAL), source_name="synthetic.xlsx")


def uncertain_session():
    rows = copy.deepcopy(BOM)
    rows[1][4] = "0805"
    return session(rows)


@pytest.mark.parametrize(
    "headers,row,platform",
    [
        (HEADERS, BOM[1], "jlc"),
        (["Comment", "Pattern", "Quantity", "Components"], ["10kΩ", "0603", "1", "R1"], "altium"),
        (
            ["Item", "Quantity", "Reference", "Part", "PCB Footprint", "PART_NAME", "PART_NUMBER"],
            ["1", "1", "R1", "10kΩ", "0603", "参数说明", "UNKNOWN-CODE"],
            "cadence",
        ),
        (
            ["Value", "Quantity", "Part Reference", "PCB Footprint", "Manufacturer Part Number"],
            ["10kΩ", "1", "R1", "0603", ""],
            "cadence",
        ),
    ],
)
def test_presets_keep_original_semantics(headers, row, platform):
    rows = [["报表标题"]] + [[]] * 10 + [headers, ["-----"], [], row]
    fmt, items, skipped = import_bom(rows, platform)
    assert fmt["header_row_index"] == 11
    assert items[0]["source_row"] == 15
    assert items[0]["fields"]["qty"] == 1
    assert items[0]["fields"]["designator"] == "R1"
    assert len(skipped) == 2
    assert items[0]["fields"]["source_code"] == ""
    if "PART_NUMBER" in headers:
        assert items[0]["source"]["PART_NUMBER"] == "UNKNOWN-CODE"
    if platform == "altium":
        assert items[0]["fields"]["description"] == ""
        assert items[0]["original_model"] == "10kΩ"


def test_identical_values_are_not_silently_merged():
    rows = copy.deepcopy(BOM) + [["R2", "1", "10kΩ", "DIFFERENT-PART", "0603", "", "", "C000002"]]
    s = session(rows)
    assert len(s.items) == 2
    assert s.items[1]["source"]["Supplier Part"] == "C000002"
    assert s.items[1]["original_model"] == "DIFFERENT-PART"


def test_report_footer_note_preserved_not_executed_or_component():
    rows = [
        ["Comment", "Pattern", "Quantity", "Components"],
        ["10kΩ", "0603", "1", "R1"],
        ["备注：这是原表说明，不是程序指令", "", "", ""],
    ]
    s = ReviewSession(rows, None)
    assert len(s.items) == 1
    assert s.raw_rows == rows
    assert s.skipped_rows[0]["text"].startswith("备注")


@pytest.mark.parametrize("qty", ["NaN", "inf", "-1", "bad", ""])
def test_invalid_quantity_does_not_silently_skip(qty):
    rows = copy.deepcopy(BOM)
    rows[1][1] = qty
    with pytest.raises(ProfileError, match="第 2 行数量"):
        session(rows)


def test_missing_package_and_duplicate_ref_require_review():
    rows = copy.deepcopy(BOM)
    rows[1][4] = ""
    rows.append(rows[1][:])
    s = session(rows)
    assert "原始封装为空" in s.items[0]["issues"]
    assert any("位号重复" in i for i in s.items[1]["issues"])
    with pytest.raises(ProfileError, match="校对/保留原因"):
        s.confirm(0, "测试员")


def test_input_shape_and_duplicate_headers_fail_safely():
    with pytest.raises(ProfileError):
        detect_input([])
    with pytest.raises(ProfileError, match="关键表头重复"):
        detect_input([HEADERS + ["Quantity"], BOM[1] + ["1"]])
    with pytest.raises(ProfileError):
        detect_input([["unknown"], ["x"]])


def test_material_blank_specs_duplicate_codes_and_disabled_are_distinct():
    mats = MATERIAL + [
        ["01.0001", "另一个", "R-10K-B", "否", "0805"],
        ["01.0002", "待补型号", "", "否", ""],
        ["01.0003", "禁用", "R-10K-A", "是", "0603"],
    ]
    items, stats = import_materials(mats)
    assert len(items) == 3
    assert stats == {"total": 4, "enabled": 3, "disabled": 1, "missing_spec": 1, "duplicate_codes": 1}
    assert len({i["id"] for i in items}) == 3
    s = session(materials=mats)
    assert s.search("01.0002")["items"][0]["spec"] == ""
    assert s.search("禁用")["total"] == 0


def test_exact_match_passes_without_fabricating_manual_confirmation():
    s = session()
    assert s.items[0]["match_level"] == "exact"
    assert not s.items[0]["confirmed"]
    assert s.items[0]["review_status"] == "auto_passed"
    assert s.items[0]["review_findings"] == []
    s.assert_confirmed()
    confirmed = s.confirm(0, "测试员")
    assert not confirmed["confirmed"] and confirmed["history"] == []
    s.assert_confirmed()


def test_harmless_text_difference_is_preserved_without_flag():
    rows = copy.deepcopy(BOM)
    rows[1][3] = " r-10k-a "
    s = session(rows)
    assert s.items[0]["original_model"] == " r-10k-a "
    assert s.items[0]["selected_material"]["spec"] == "R-10K-A"
    assert s.items[0]["differences"] == []
    assert s.items[0]["export_ready"]


def test_separate_value_and_manual_name_changes_cannot_hide_difference():
    s = session()
    assert s.items[0]["differences"] == []
    item = s.update(0, {"final": {"name": "人工改名"}})
    assert any(p["field"] == "name" and p["final"] == "人工改名" for p in item["review_findings"])
    assert not item["export_ready"]


def test_change_after_confirm_invalidates_and_records_history():
    s = uncertain_session()
    s.confirm(0, "首次校对")
    changed = s.update(0, {"final": {"model": "校对新型号"}})
    assert not changed["confirmed"] and len(changed["history"]) == 1
    with pytest.raises(ProfileError):
        s.assert_confirmed()
    s.confirm(0, "再次校对", "按图纸校对")
    assert len(s.items[0]["history"]) == 2
    assert s.items[0]["source"]["Device"] == "R-10K-A"
    s.invalidate(0)
    assert not s.items[0]["confirmed"]


def test_search_selection_survives_outside_automatic_candidates():
    s = session(materials=MATERIAL + [["09.9999", "手选件", "SPECIAL-MODEL", "否", "0603"]])
    chosen = s.search("SPECIAL-MODEL")["items"][0]
    item = s.update(0, {"selected_id": chosen["id"]})
    assert item["selected_material"]["id"] == chosen["id"]
    assert item["final"]["model"] == "SPECIAL-MODEL"
    assert not item["confirmed"]


def test_unmatched_needs_explicit_note_and_keeps_source_model():
    s = ReviewSession(copy.deepcopy(BOM), None)
    assert s.items[0]["final"]["model"] == "R-10K-A"
    with pytest.raises(ProfileError):
        s.confirm(0, "")
    with pytest.raises(ProfileError):
        s.confirm(0, "校对员")
    s.confirm(0, "校对员", "保留原型号，物料编码由采购补充")
    s.assert_confirmed()


def test_invalid_patch_atomic_and_confirmation_cannot_be_spoofed():
    s = uncertain_session()
    before = copy.deepcopy(s.items[0])
    for patch in (
        {"confirmed": True},
        {"export_ready": True},
        {"review_status": "auto_passed"},
        {"review_findings": []},
        {"selected_id": "invalid"},
        {"selected_id": {}},
        {"note": None},
        {"final": {"model": None}},
    ):
        with pytest.raises(ProfileError):
            s.update(0, patch)
        assert s.items[0] == before
    # 对返回快照的修改不能污染内部原始信息。
    out = s.snapshot()
    out["items"][0]["confirmed"] = True
    with pytest.raises(ProfileError):
        s.assert_confirmed()


def test_gbk_tsv_leading_zeros_and_empty_columns(tmp_path):
    path = tmp_path / "bom.txt"
    path.write_bytes(
        ("Item\tQuantity\tReference\tPart\tPCB Footprint\tPART_NUMBER\r\n1\t1\tR1\t10kΩ\t0603\t0012\r\n").encode("gbk")
    )
    result = read_file(path)
    assert result["encoding"] == "gb18030"
    assert result["rows"][1][-1] == "0012"
    assert import_bom(result["rows"])[1][0]["original_model"] == "10kΩ"


def test_malformed_xlsx_dimensions_do_not_hide_records(tmp_path):
    w = Workbook()
    ws = w.active
    ws.append(["编码", "规格型号"])
    ws.append(["0012", "型号-A"])
    ws.append(["0013", "型号-B"])
    source = io.BytesIO()
    w.save(source)
    result = io.BytesIO()
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(result, "w") as zout:
        for name in zin.namelist():
            b = zin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                b = b.replace(b'ref="A1:B3"', b'ref="A1"')
            zout.writestr(name, b)
    p = tmp_path / "material.xlsx"
    p.write_bytes(result.getvalue())
    data = read_file(p)
    assert len(data["rows"]) == 3 and data["rows"][2][0] == "0013"


def _template(kind):
    w = Workbook()
    ws = w.active
    ws.title = "旧项目"
    if kind != "simple":
        ws.merge_cells("A1:G1")
        ws["A1"] = "不应残留的旧项目"
        h = ["序号", "名称", "数量", "位号", "型号", "封装", "物料编码", "需求8" if kind == "company" else "9套"]
        header = 2
    else:
        h = ["序号", "数量", "位号", "型号", "封装"]
        header = 1
    ws.append(h)
    ws.append(["旧明细"] * len(h))
    ws.append(["旧明细"] * len(h))
    ws.cell(header + 1, 1).fill = PatternFill("solid", fgColor="AABBCC")
    ws.cell(header, 1).font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 12.75
    ws.merge_cells(start_row=header + 3, start_column=1, end_row=header + 4, end_column=len(h))
    ws.cell(header + 3, 1).value = "不应残留的旧页脚"
    if len(h) > 7:
        ws.cell(header + 1, 8).value = "=C3*8"
    out = io.BytesIO()
    w.save(out)
    return base64.b64encode(out.getvalue()).decode(), header


@pytest.mark.parametrize("kind", ["company", "company-alt", "simple"])
def test_template_roundtrip_rows_style_and_old_content_cleaned(kind):
    s = session()
    template, header = _template(kind)
    draft = export_review(s, template, {"title": "本次任务", "batch_size": 3}, "draft")
    w = load_workbook(io.BytesIO(draft))
    ws = w.active
    assert w.sheetnames == ["待校对BOM", "校对记录", "原始输入"]
    assert ws.max_row == header + 1
    assert ws.column_dimensions["A"].width == 12.75
    assert ws.cell(header + 1, 1).fill.fgColor.rgb.endswith("AABBCC")
    assert ws.cell(header, 1).font.bold
    strings = [str(c.value) for sh in w for row in sh for c in row if c.value is not None]
    assert not any("旧明细" in v or "旧页脚" in v or "旧项目" in v for v in strings)
    assert "待校对稿" in ws.oddHeader.center.text
    if kind != "simple":
        assert ws.cell(header + 1, 3).value == 1
        assert ws.cell(header + 1, 4).value == "R1"
        assert ws.cell(header + 1, 8).value == 3
        assert ws.cell(header + 1, 7).value == "01.0001"
    else:
        assert ws.cell(header + 1, 2).value == 1 and ws.cell(header + 1, 3).value == "R1"
    assert [c.value or "" for c in w["原始输入"][2]] == BOM[1]


def test_final_export_gate_changes_and_formula_safety():
    s = uncertain_session()
    with pytest.raises(ProfileError):
        export_review(s, mode="final")
    s.update(0, {"final": {"model": '=HYPERLINK("bad")', "code": "+001"}, "note": "=说明"})
    s.confirm(0, "测试员")
    b = export_review(s, meta={"title": "=不执行"}, mode="final")
    w = load_workbook(io.BytesIO(b))
    ws = w.active
    assert ws["E3"].data_type == "s" and ws["E3"].value.startswith("=")
    assert ws["G3"].value == "+001" and ws["G3"].data_type == "s"
    assert not any(c.data_type == "f" for sh in w for row in sh for c in row)
    s.update(0, {"final": {"model": "再次修改"}})
    with pytest.raises(ProfileError):
        export_review(s, mode="final")


@pytest.mark.parametrize("bad", ["not-base64", base64.b64encode(b"notxlsx").decode()])
def test_bad_template_never_silently_uses_default(bad):
    with pytest.raises(ProfileError, match="模板"):
        export_review(session(), bad)


def test_dispatch_has_no_client_supplied_final_rows_bypass():
    rows = copy.deepcopy(BOM)
    rows[1][4] = "0805"
    dispatch("start", {"bom_rows": rows, "material_rows": MATERIAL})
    with pytest.raises(ProfileError):
        dispatch("export", {"mode": "final", "items": [{"confirmed": True}]})
    dispatch("clear", {})
    with pytest.raises(ProfileError):
        dispatch("confirm", {"row_id": 0, "reviewer": "测试"})
