"""显式私有回归矩阵：不打包/不上传原文件；技术模拟确认绝非物料业务验收。

在仓库根运行：.venv/bin/python core/tests/private_regression.py
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import time
from pathlib import Path

from openpyxl import load_workbook

from bomcore.review_api import ReviewSession
from bomcore.review_export import export_review
from bomcore.review_import import read_file
from bomcore.schema import ProfileError

ROOT = Path(__file__).resolve().parents[2]
PRIVATE = Path(__file__).parent / "fixtures/private/closed-loop"
INPUTS = PRIVATE / "inputs"
OUTPUTS = PRIVATE / "outputs"
CASES = {
    "jlc.xlsx": (55, 171),
    "altium.xlsx": (94, 250),
    "cadence-report.xlsx": (309, 1004),
    "cadence-detail.xlsx": (76, 76),
    "cadence-grouped.xlsx": (33, 76),
    "cadence-grouped.txt": (33, 76),
}
TEMPLATES = ["template-company.xlsx", "template-company-alt.xlsx", "template-simple.xlsx"]


def verify(data: bytes, session: ReviewSession, template: str, mode: str) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=False)
    ws = wb.active
    header = 1 if template == "template-simple.xlsx" else 2
    cols = {c.value: c.column for c in ws[header] if c.value}
    assert ws.max_row - header == len(session.items), "输出行数改变"
    assert "校对记录" in wb.sheetnames and "原始输入" in wb.sheetnames
    assert wb["校对记录"].max_row == len(session.items) + 1
    assert not any(c.data_type == "f" for sh in wb for row in sh for c in row), "意外公式"
    for offset, item in enumerate(session.items, header + 1):
        for name, expected in (
            ("数量", item["fields"]["qty"]),
            ("位号", item["fields"]["designator"]),
            ("型号", item["final"]["model"]),
            ("封装", item["final"]["footprint"]),
        ):
            actual = ws.cell(offset, cols[name]).value
            assert (actual if actual is not None else "") == expected, f"字段{name}不一致，源行{item['source_row']}"
        if "物料编码" in cols:
            assert (ws.cell(offset, cols["物料编码"]).value or "") == item["final"]["code"]
        state = ws.cell(offset, cols["校对状态"]).value
        expected_state = "已人工确认" if item["confirmed"] else (None if item["export_ready"] else "待校对（不可投产）")
        assert state == expected_state
    source = wb["原始输入"]
    for r, row in enumerate(session.raw_rows, 1):
        for c, value in enumerate(row, 1):
            assert (source.cell(r, c).value or "") == value, "原始文本丢失或被更改"
    original = load_workbook(INPUTS / template)
    assert ws.column_dimensions["A"].width == original.active.column_dimensions["A"].width
    for region in ws.merged_cells.ranges:
        assert region.max_row <= header, "旧数据区合并残留"
    original.close()
    quantity = sum(ws.cell(r, cols["数量"]).value for r in range(header + 1, ws.max_row + 1))
    wb.close()
    return {
        "rows": len(session.items),
        "quantity": quantity,
        "bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "checks": ["行数", "数量", "位号", "型号", "编码", "原文", "样式", "公式安全", "确认状态"],
    }


def main():
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    material = read_file(INPUTS / "material.xlsx")
    report = {"warning": "技术自动化测试；任何模拟确认不代表真实物料已审核，输出禁止投产", "cases": []}
    for name, (expected_rows, expected_quantity) in CASES.items():
        started = time.monotonic()
        payload = read_file(INPUTS / name)
        s = ReviewSession(payload["rows"], material["rows"], source_name=name, sheet_name=payload["sheet_name"])
        assert len(s.items) == expected_rows
        assert sum(i["fields"]["qty"] for i in s.items) == expected_quantity
        assert s.material_stats["total"] == 17371
        initial_stats = s.snapshot()["stats"]
        blocked = False
        try:
            export_review(s, mode="final")
        except ProfileError as e:
            blocked = e.code == "CONFIRMATION_REQUIRED"
        assert blocked == (initial_stats["pending"] > 0), "正式导出门禁与未处理问题不一致"
        checks = []
        for template in TEMPLATES:
            encoded = base64.b64encode((INPUTS / template).read_bytes()).decode()
            data = export_review(s, encoded, {"title": f"技术测试·待校对·{Path(name).stem}"}, "draft")
            filename = f"{Path(name).stem}-{Path(name).suffix[1:]}-{Path(template).stem}-待校对.xlsx"
            (OUTPUTS / filename).write_bytes(data)
            checks.append(
                {"template": template, "mode": "draft", "file": filename, **verify(data, s, template, "draft")}
            )
        # 仅验证技术状态机：明确保留原文，模拟校对人，输出名称与标题必须包含不可投产。
        for item in list(s.items):
            s.update(
                item["row_id"],
                {"selected_id": None, "note": "自动化技术测试：仅保留原始型号，不确认企业物料，禁止投产"},
            )
            s.confirm(item["row_id"], "自动化模拟确认（非业务审核）")
        for template in TEMPLATES:
            encoded = base64.b64encode((INPUTS / template).read_bytes()).decode()
            data = export_review(s, encoded, {"title": f"技术模拟·禁止投产·{Path(name).stem}"}, "final")
            filename = f"{Path(name).stem}-{Path(name).suffix[1:]}-{Path(template).stem}-技术模拟禁止投产.xlsx"
            (OUTPUTS / filename).write_bytes(data)
            checks.append(
                {
                    "template": template,
                    "mode": "simulated-final",
                    "file": filename,
                    **verify(data, s, template, "final"),
                }
            )
        s.update(0, {"final": {"model": s.items[0]["original_model"] + " [模拟修改]"}})
        try:
            export_review(s, mode="final")
            raise AssertionError("修改后仍允许正式导出")
        except ProfileError as e:
            assert e.code == "CONFIRMATION_REQUIRED"
        report["cases"].append(
            {
                "input": name,
                "format": s.profile["id"],
                "rows": expected_rows,
                "quantity": expected_quantity,
                "elapsed_s": round(time.monotonic() - started, 3),
                "unconfirmed_blocked": True,
                "initial_stats": initial_stats,
                "edited_confirmation_invalidated": True,
                "outputs": checks,
            }
        )
        print(name, "rows", expected_rows, "quantity", expected_quantity, "6 template/mode checks passed")
    (OUTPUTS / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        "Private matrix complete: 6 inputs x 3 templates x 2 modes; details kept in ignored outputs/verification.json"
    )


if __name__ == "__main__":
    main()
