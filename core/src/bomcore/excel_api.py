"""Excel 优先：保留旧公司分组、候选展开及主表，不设网页审批门禁。"""

from __future__ import annotations

import copy
import io
import json
from collections import Counter, defaultdict
from dataclasses import asdict
from importlib import resources

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from .grouping import clean_secondary_category, group_and_sort
from .matching import build_material_rc_index, match_device
from .models import BomItem, MaterialItem, cell_text, is_dnp
from .render import (
    FILL_DNP,
    FILL_MATCH,
    FILL_MISSING,
    FILL_MULTI_A,
    FILL_MULTI_B,
    FILL_MULTI_PARAM_A,
    FILL_MULTI_PARAM_B,
    FILL_NON_COMPONENT,
    FILL_NONE,
    FILL_PARAM,
    _resolve_row,
    render,
)
from .review_export import safe_cell
from .review_import import import_bom, normalize_header, validate_rows
from .review_rules import candidate_findings, material_evidence
from .schema import ProfileError

AMBER = PatternFill("solid", fgColor="FFEB9C", bgColor="FFEB9C")
STATUS_FILLS = {
    "exact": FILL_MATCH,
    "model": FILL_MATCH,
    "substring": AMBER,
    "param": FILL_PARAM,
    "none": FILL_NONE,
    "non_component": FILL_NON_COMPONENT,
}


def material_input(rows):
    """金蝶完整表，或旧格式的规格→编码/金蝶系统型号字典。"""
    if rows is None:
        return "none", [], {}
    validate_rows(rows)
    for index, row in enumerate(rows[:50]):
        headers = [normalize_header(h) for h in row]

        def column(*names, headers=headers):
            return next((c for c, h in enumerate(headers) if h in names), None)

        code_col = column("编码", "物料编码", "code")
        spec_col = column("规格型号", "spec")
        kind = "kingdee"
        if spec_col is None:
            spec_col = column("规格", "device")
            kind = "legacy"
        if code_col is None or spec_col is None:
            continue
        name_col = column("名称", "物料名称", "name")
        disabled_col = column("禁用状态", "disabled", "status")
        system_col = column("金蝶系统型号", "系统型号")
        materials, lookup = [], {}
        for raw in rows[index + 1 :]:

            def get(col, raw=raw):
                return cell_text(raw[col]) if col is not None and col < len(raw) else ""

            if get(disabled_col).lower() in ("是", "disabled", "true", "1", "yes"):
                continue
            spec, code = get(spec_col), get(code_col)
            if not spec:
                continue
            if kind == "legacy":
                lookup[spec] = {"code": code, "name": "", "spec": get(system_col)}
            else:
                materials.append(MaterialItem(code=code, name=get(name_col), spec=spec))
        return kind, materials, lookup
    raise ProfileError("INVALID_PROFILE", "物料表需包含编码/规格型号，或旧格式的物料编码/规格/金蝶系统型号")


def _empty_match(level):
    return {
        "level": level,
        "status_text": "" if level == "skipped" else "未匹配",
        "confidence": None,
        "candidates": [],
        "selected": 0,
        "manual_code": None,
    }


def analyze_excel(bom_rows, material_rows=None, platform="auto"):
    profile, source, skipped = import_bom(bom_rows, platform)
    kind, materials, lookup = material_input(material_rows)
    raw_groups = defaultdict(list)
    inputs = []
    for row in source:
        f = row["fields"]
        item = BomItem(
            **{
                k: f[k]
                for k in (
                    "value",
                    "designator",
                    "qty",
                    "footprint",
                    "mpn",
                    "manufacturer",
                    "description",
                    "category",
                    "tolerance",
                )
            },
            dnp=is_dnp(f["description"]),
        )
        if profile["platform"] != "jlc" and not item.mpn.strip():
            item.mpn = item.value
        inputs.append(item)
        raw_groups[(item.value, item.footprint, item.dnp)].append(row)
    merged = group_and_sort(inputs)
    rc_index = build_material_rc_index(materials)
    items, stats, cache = [], Counter(), {}
    for row_id, item in enumerate(merged):
        fields = asdict(item)
        device = item.mpn
        key = (device, item.value, item.footprint, item.tolerance)
        if key not in cache:
            if kind == "none":
                result = _empty_match("skipped")
            elif not device:
                result = _empty_match("none")
            elif kind == "legacy":
                candidate = lookup.get(device.strip())
                result = _empty_match("none")
                if candidate and (candidate["code"] or candidate["spec"]):
                    result.update(level="exact", status_text="精确匹配", confidence="high", candidates=[candidate])
            else:
                result = match_device(
                    device, materials, rc_index, name=item.value, footprint=item.footprint, tolerance=item.tolerance
                )
            cache[key] = result
        match = copy.deepcopy(cache[key])
        match["selected"] = next((i for i, c in enumerate(match["candidates"]) if c["code"].startswith("01.")), 0)
        match["manual_code"] = None
        stats[match["level"]] += 1
        fields["category"] = clean_secondary_category(fields["category"])
        items.append(
            {
                "row_id": row_id,
                "fields": fields,
                "match": match,
                "_sources": raw_groups[(item.value, item.footprint, item.dnp)],
            }
        )
    return {
        "items": items,
        "profile": profile,
        "skipped": skipped,
        "material_kind": kind,
        "stats": {
            "source_rows": len(source),
            "groups": len(items),
            "quantity": sum(i.qty for i in inputs),
            "material_count": len(lookup) if kind == "legacy" else len(materials),
            "multi": stats["multi"],
            "unmatched": stats["none"],
            "non_component": stats["non_component"],
            "low_confidence": stats["substring"],
        },
    }


def _company_workbook(items, meta, has_material):
    for item in items:
        fields = item["fields"]
        # Excel单元格上限不是业务行数上限。只在必要时去除位号分隔空格，避免静默截断。
        if len(fields["designator"]) > 32767:
            fields["designator"] = fields["designator"].replace(", ", ",")
        if any(isinstance(v, str) and len(v) > 32767 for v in fields.values()):
            raise ProfileError("INVALID_ROWS", "合并后字段超过Excel单元格32767字符限制，请拆分输入后重试")
    profile = json.loads(resources.files("bomcore").joinpath("presets/default_output_template.json").read_text())
    if not has_material:
        profile["columns"].pop("L")
    workbook = load_workbook(io.BytesIO(render(items, profile, meta)))
    ws = workbook.active
    row_number, multi_group, quantity = 5, 0, 0
    for item in items:
        fields = item["fields"]
        resolved = _resolve_row(item)
        group_fill = None
        if resolved["expand"]:
            a, b = (
                (FILL_MULTI_PARAM_A, FILL_MULTI_PARAM_B)
                if item["match"]["status_text"].startswith("参数匹配")
                else (FILL_MULTI_A, FILL_MULTI_B)
            )
            group_fill = a if multi_group % 2 == 0 else b
            multi_group += 1
        for result in resolved["rows"]:
            # 旧公司备注列是未填写单元格，不写空字符串占位。
            ws.cell(row_number, 9).value = None
            ws.cell(row_number, 9).data_type = "n"
            # 按候选组交替，A-I铺色；单匹配仅状态列着色，恢复旧辅助列及DNP颜色。
            for cell in ws[row_number]:
                cell.fill = PatternFill()
            if group_fill:
                for col in range(1, 10):
                    ws.cell(row_number, col).fill = group_fill
            if has_material:
                fill = group_fill or STATUS_FILLS.get(result["level"])
                if fill:
                    ws.cell(row_number, 12).fill = fill
            if has_material and fields["mpn"] and result["level"] not in ("non_component", "skipped"):
                if not result.get("code"):
                    ws.cell(row_number, 2).fill = FILL_MISSING
                if not result.get("spec"):
                    ws.cell(row_number, 4).fill = FILL_MISSING
            if is_dnp(fields["description"]):
                ws.cell(row_number, 11).fill = FILL_DNP
            quantity += fields["qty"]
            row_number += 1
    return workbook, row_number - 5, quantity


_CONFLICTS = {
    "value_mismatch",
    "package_mismatch",
    "property_mismatch",
    "manufacturer_mismatch",
    "source_value_conflict",
    "source_property_conflict",
    "source_package_conflict",
    "library_package_conflict",
}


def _trace(workbook, analysis, bom_rows):
    raw = workbook.create_sheet("原始输入")
    for nr, row in enumerate(bom_rows, 1):
        for nc, value in enumerate(row, 1):
            safe_cell(raw, nr, nc, value)
    raw.freeze_panes = f"A{analysis['profile']['header_row_index'] + 2}"
    notes = workbook.create_sheet("Excel校对提示")
    for c, title in enumerate(("主表序号", "位号", "请在Excel中处理（不阻断导出）"), 1):
        safe_cell(notes, 1, c, title).font = Font(bold=True)
    notes.column_dimensions["A"].width = 14
    notes.column_dimensions["B"].width = 35
    notes.column_dimensions["C"].width = 100
    nr = 2
    for item in analysis["items"]:
        match, fields = item["match"], item["fields"]
        messages = []
        if match["level"] == "multi":
            messages.append("同组候选已展开；请在Excel保留正确候选、删除其余候选，不能直接将展示数量合计用于采购")
        elif match["level"] == "substring":
            messages.append("低置信度料号匹配，请在Excel核对完整规格；主表状态列按旧工具标为琥珀色")
        elif match["level"] == "param":
            messages.append("参数匹配，请在Excel核对电压、介质、厂商等完整要求")
        elif match["level"] == "none":
            messages.append("未匹配，请在Excel补充或核实物料编码/规格")
        sources = item["_sources"]
        for key, label in (("mpn", "型号"), ("manufacturer", "厂商"), ("tolerance", "精度")):
            values = {r["fields"][key].strip() for r in sources if r["fields"][key].strip()}
            if len(values) > 1:
                messages.append(
                    f"组内{label}不同；主表仍按旧规则合并并取首个非空值，原始输入附表保留全部原值，请在Excel核对是否需拆分"
                )
        if match["level"] != "non_component":
            for source in sources:
                messages.extend(source["issues"])
            if match["candidates"]:
                chosen = match["candidates"][match["selected"]]
                material = {**chosen, "footprint": "", "tolerance": "", "manufacturer": ""}
                evidence = material_evidence(material)
                for source in sources:
                    findings, _ = candidate_findings(
                        source, material, evidence, analysis["profile"]["platform"] == "jlc"
                    )
                    messages.extend(
                        f"{f['reason']}；原值：{f['original']}；库值：{f['library']}"
                        for f in findings
                        if f["code"] in _CONFLICTS
                    )
        if messages:
            text = "；".join(dict.fromkeys(messages))
            if len(text) > 32000:
                text = text[:32000] + "…提示过长，请结合完整原始输入核对"
            for c, value in enumerate((item["row_id"] + 2, fields["designator"], text), 1):
                safe_cell(notes, nr, c, value)
            nr += 1
    if nr == 2:
        safe_cell(notes, nr, 3, "没有额外提示。此文件为转换结果，不代表物料已经工程审批。")
    notes.freeze_panes = "C2"


def convert_excel(args: dict) -> dict:
    """无确认会话的一键转换；Excel编辑是用户明确选择的主工作流。"""
    if not isinstance(args, dict):
        raise ProfileError("INVALID_ROWS", "转换参数必须为对象")
    meta = args.get("meta") or {}
    if not isinstance(meta, dict) or any(not isinstance(v, str) or len(v) > 32767 for v in meta.values()):
        raise ProfileError("INVALID_ROWS", "表头信息必须为有效文本")
    include_trace = args.get("include_trace", False)
    if not isinstance(include_trace, bool):
        raise ProfileError("INVALID_ROWS", "附表选项必须为布尔值")
    analysis = analyze_excel(args.get("bom_rows"), args.get("material_rows"), args.get("platform", "auto"))
    workbook, count, quantity = _company_workbook(analysis["items"], meta, args.get("material_rows") is not None)
    if include_trace:
        _trace(workbook, analysis, args["bom_rows"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return {
        "data": output.getvalue(),
        "stats": {**analysis["stats"], "output_rows": count, "display_quantity": quantity},
        "format": analysis["profile"]["name"],
        "material_format": {"none": "未提供", "kingdee": "金蝶完整物料表", "legacy": "旧简易物料表"}[
            analysis["material_kind"]
        ],
    }
