"""v2 输入：固定 EDA 预设、原文留存、TXT 编码兼容、逐行而非静默合并。"""

from __future__ import annotations

import csv
import io
import json
import math
import re
import zipfile
from importlib import resources
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format

from .detect import normalize_header
from .models import is_dnp
from .schema import ProfileError

MAX_ROWS = 100_000
MAX_COLUMNS = 256
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 120 * 1024 * 1024
TEXT_FIELDS = (
    "designator",
    "value",
    "mpn",
    "footprint",
    "manufacturer",
    "description",
    "category",
    "tolerance",
    "source_code",
)


def _text(value) -> str:
    return "" if value is None else str(value)


def validate_rows(rows) -> None:
    if not isinstance(rows, list) or not rows or len(rows) > MAX_ROWS:
        raise ProfileError("INVALID_ROWS", f"表格为空或超过 {MAX_ROWS} 行限制")
    for row in rows:
        if not isinstance(row, list) or len(row) > MAX_COLUMNS:
            raise ProfileError("INVALID_ROWS", f"表格格式错误或超过 {MAX_COLUMNS} 列限制")
        if any(not isinstance(v, str) or len(v) > 32767 for v in row):
            raise ProfileError("INVALID_ROWS", "单元格必须是文本且不超过 32767 字符")


def check_xlsx_bytes(data: bytes) -> None:
    if len(data) > MAX_FILE_BYTES:
        raise ProfileError("INVALID_ROWS", "文件超过 20MB，请拆分后重试")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ProfileError("INVALID_ROWS", "表格解压后体积过大")
    except zipfile.BadZipFile as exc:
        raise ProfileError("INVALID_ROWS", "不是有效的 XLSX 文件") from exc


def read_file(path: str | Path, sheet_name: str | None = None) -> dict:
    """CLI/私有回归读取器；不执行宏、链接或单元格公式。"""
    path = Path(path)
    data = path.read_bytes()
    if len(data) > MAX_FILE_BYTES:
        raise ProfileError("INVALID_ROWS", "文件超过 20MB，请拆分后重试")
    if path.suffix.lower() in (".txt", ".tsv", ".csv"):
        encoding = "utf-8-sig"
        try:
            content = data.decode(encoding)
        except UnicodeDecodeError:
            encoding = "gb18030"
            try:
                content = data.decode(encoding)
            except UnicodeDecodeError as exc:
                raise ProfileError("INVALID_ROWS", "文本编码无法识别，请转存 UTF-8 后导入") from exc
        delimiter = "\t" if "\t" in content else ","
        rows = list(csv.reader(io.StringIO(content), delimiter=delimiter))
        validate_rows(rows)
        return {"rows": rows, "sheet_name": path.stem, "encoding": encoding}
    if path.suffix.lower() != ".xlsx":
        raise ProfileError("INVALID_ROWS", "支持 XLSX、TXT、TSV、CSV；请先将 XLS 转存 XLSX")
    check_xlsx_bytes(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            if sheet_name is not None and sheet_name not in wb.sheetnames:
                raise ProfileError("INVALID_ROWS", "所选工作表不存在")
            sheets = [wb[sheet_name]] if sheet_name else wb.worksheets
            for ws in sheets:
                # 一些 ERP 导出错误声明 dimension=A1，不能用声明范围截断真实数据。
                ws.reset_dimensions()
                rows = []
                for row in ws.iter_rows():
                    values = []
                    for cell in row:
                        v = cell.value
                        fmt = getattr(cell, "number_format", "General")
                        if isinstance(v, (int, float)) and re.fullmatch(r"0+", fmt):
                            values.append(str(int(v)).zfill(len(fmt)) if float(v).is_integer() else str(v))
                        elif v is not None and is_date_format(fmt) and hasattr(v, "isoformat"):
                            values.append(v.isoformat(sep=" "))
                        elif isinstance(v, float) and v.is_integer():
                            values.append(str(int(v)))
                        else:
                            values.append(_text(v))
                    rows.append(values)
                    if len(rows) > MAX_ROWS:
                        raise ProfileError("INVALID_ROWS", "工作表行数超过限制")
                if any(any(v for v in r) for r in rows):
                    validate_rows(rows)
                    return {"rows": rows, "sheet_name": ws.title}
        finally:
            wb.close()
    except ProfileError:
        raise
    except Exception as exc:
        raise ProfileError("INVALID_ROWS", "无法读取工作表，请检查文件是否损坏") from exc
    raise ProfileError("INVALID_ROWS", "工作表为空")


def eda_presets() -> list[dict]:
    return json.loads(resources.files("bomcore").joinpath("presets/eda_inputs_v2.json").read_text(encoding="utf-8"))


def detect_input(rows: list[list[str]], platform: str = "auto") -> dict:
    validate_rows(rows)
    presets = [p for p in eda_presets() if platform == "auto" or platform in (p["platform"], p["id"])]
    for index, row in enumerate(rows[:50]):
        headers = [normalize_header(h) for h in row]
        for preset in presets:
            if set(preset["required_headers"]).issubset(headers):
                required = preset["required_headers"]
                if any(headers.count(h) != 1 for h in required):
                    raise ProfileError("INVALID_ROWS", "关键表头重复，请先校对原表")
                return {
                    "id": preset["id"],
                    "platform": preset["platform"],
                    "name": preset["name"],
                    "header_row_index": index,
                    "column_map": {
                        row[i]: preset["column_map"][h] for i, h in enumerate(headers) if h in preset["column_map"]
                    },
                }
    raise ProfileError("INVALID_PROFILE", "未识别到已支持的 EDA 表头。请确认平台、工作表或导出格式")


def references(text: str) -> list[str]:
    """仅计数/检查，不改写位号原文，不猜范围表达式。"""
    return [r for r in re.split(r"[,，;；\s]+", text.strip()) if r]


def import_bom(rows: list[list[str]], platform: str = "auto") -> tuple[dict, list[dict], list[dict]]:
    profile = detect_input(rows, platform)
    header_index = profile["header_row_index"]
    header = rows[header_index]
    mapping = profile["column_map"]
    items, skipped = [], []
    seen_refs: dict[str, int] = {}
    for index, raw in enumerate(rows[header_index + 1 :], header_index + 1):
        if not any(v.strip() for v in raw):
            skipped.append({"row": index + 1, "reason": "空行"})
            continue
        if all(not v.strip() or re.fullmatch(r"[-=_ ]+", v.strip()) for v in raw):
            skipped.append({"row": index + 1, "reason": "报表分隔线"})
            continue
        fields = dict.fromkeys(TEXT_FIELDS, "")
        source = {}
        qty_text = ""
        for col, name in enumerate(header):
            value = raw[col] if col < len(raw) else ""
            key = name or f"未命名列{col + 1}"
            if key in source:
                key = f"{key}（列{col + 1}）"
            source[key] = value
            field = mapping.get(name)
            if field == "qty":
                qty_text = value
            elif field in fields and not fields[field]:
                fields[field] = value
        # 已知 Altium 导出可能在明细后附整行备注。保留但不将其当元件或指令。
        if (
            not qty_text.strip()
            and not fields["designator"].strip()
            and raw
            and re.match(r"^(备注|说明|注[:：]|note[:：])", raw[0].strip(), re.IGNORECASE)
        ):
            skipped.append(
                {
                    "row": index + 1,
                    "reason": "表尾备注（请核对原始输入页）",
                    "text": " ".join(v for v in raw if v.strip()),
                }
            )
            continue
        try:
            qty = float(qty_text.strip().replace(",", ""))
            if not math.isfinite(qty) or qty < 0:
                raise ValueError
        except ValueError as exc:
            raise ProfileError("INVALID_ROWS", f"第 {index + 1} 行数量不是有效的非负数，请校对原表") from exc
        fields["qty"] = int(qty) if qty.is_integer() else qty
        fields["dnp"] = is_dnp(fields["description"]) or is_dnp(fields["value"])
        model = fields["mpn"] if fields["mpn"].strip() else fields["value"]
        issues = []
        if not model.strip():
            issues.append("原始型号/数值为空")
        if not fields["footprint"].strip():
            issues.append("原始封装为空")
        refs = references(fields["designator"])
        if not refs:
            issues.append("位号为空")
        elif len(refs) != qty:
            issues.append("数量与可识别位号数不一致，请核对")
        for ref in refs:
            if ref in seen_refs:
                issues.append(f"位号重复（首次出现在第 {seen_refs[ref]} 行）")
            else:
                seen_refs[ref] = index + 1
        if fields["dnp"]:
            issues.append("原文标记 DNP，不自动删除或改数量")
        items.append(
            {
                "row_id": len(items),
                "source_row": index + 1,
                "source": source,
                "fields": fields,
                "original_model": model,
                "issues": list(dict.fromkeys(issues)),
            }
        )
    if not items:
        raise ProfileError("INVALID_ROWS", "表头下方没有 BOM 明细")
    return profile, items, skipped
